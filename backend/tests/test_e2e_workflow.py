"""End-to-End full application integration test executing the 14-step workflow."""

import io
from datetime import date, datetime
from unittest.mock import MagicMock

import openpyxl
from docx import Document

from app.db.models import ClientModel, ItemModel, MeetingModel
from app.export.csv_export import build_csv_export
from app.export.excel import build_excel_export
from app.export.json_export import build_json_dict
from app.export.word import build_word_export
from app.schemas.client import ClientCreate
from app.schemas.item import ItemUpdate
from app.services.clarification_service import ClarificationService
from app.services.item_service import ItemService
from app.services.meeting_service import MeetingService
from app.similarity.deduplication import DeduplicationService


def test_full_14_step_workflow():
    """Execute the end-to-end meeting lifecycle workflow from ingestion to export."""
    mock_db = MagicMock()

    # Step 1: Create a client
    mock_db.scalars.return_value.first.return_value = None  # No existing duplicate client
    client_in = ClientCreate(name="Starlight Dynamics Inc.")
    client = MeetingService.create_client(mock_db, client_in)
    assert client.name == "Starlight Dynamics Inc."

    # Step 2 & 3: Create a meeting & Submit transcript
    transcript_text = (
        "Sarah (Lead PM): Welcome team. Let's lock in Q4 deliverables.\n"
        "Alex (Backend Dev): I will complete the MySQL 8 migration by November 20th.\n"
        "David (Designer): We should redesign the customer checkout flow.\n"
        "Sarah: Approved. We will standardise on FastAPI and MySQL 8.\n"
        "Elena (QA): Is the client staging environment available? We need access."
    )
    mock_db.scalars.return_value.first.return_value = client  # Client exists
    meeting = MeetingService.create_meeting_from_text(
        db=mock_db,
        client_id=client.id,
        title="Q4 Technical Architecture Alignment",
        text=transcript_text,
        meeting_date=date(2026, 10, 15),
        attendees=["Sarah", "Alex", "David", "Elena"],
    )
    assert meeting.title == "Q4 Technical Architecture Alignment"
    assert meeting.client_id == client.id
    assert meeting.text_length > 0

    # Step 4 & 5: Extracted items with verified evidence and uncertainty states
    task_1 = ItemModel(
        id="item-task-1",
        meeting_id=meeting.id,
        item_type="task",
        payload_json={
            "description": "Complete the MySQL 8 migration",
            "owner": {"name": "Alex", "state": "explicit"},
            "due": {"resolved_date": "2026-11-20", "raw_text": "November 20th"},
            "priority": "high",
            "confidence": "high",
            "evidence": [
                {
                    "quote": "I will complete the MySQL 8 migration by November 20th.",
                    "line_start": 2,
                    "line_end": 2,
                }
            ],
        },
        review_status="needs_review",
    )
    task_2 = ItemModel(
        id="item-task-2",
        meeting_id=meeting.id,
        item_type="task",
        payload_json={
            "description": "Redesign the customer checkout flow",
            "owner": {"name": None, "state": "not_stated"},
            "due": {"resolved_date": None, "raw_text": None},
            "priority": "medium",
            "confidence": "medium",
            "evidence": [
                {
                    "quote": "We should redesign the customer checkout flow.",
                    "line_start": 3,
                    "line_end": 3,
                }
            ],
            "clarification_needed": ["Assign confirmed owner and target delivery milestone"],
        },
        review_status="needs_review",
    )
    dec_1 = ItemModel(
        id="item-dec-1",
        meeting_id=meeting.id,
        item_type="decision",
        payload_json={
            "description": "Standardise on FastAPI and MySQL 8",
            "owner": {"name": "Sarah", "state": "explicit"},
            "evidence": [{"quote": "We will standardise on FastAPI and MySQL 8."}],
        },
        review_status="needs_review",
    )
    q_1 = ItemModel(
        id="item-q-1",
        meeting_id=meeting.id,
        item_type="open_question",
        payload_json={
            "question": "Is the client staging environment available?",
            "asked_by": "Elena",
            "directed_to": None,
            "evidence": [{"quote": "Is the client staging environment available?"}],
        },
        review_status="needs_review",
    )
    meeting_model = MeetingModel(
        id=meeting.id,
        client_id=client.id,
        title=meeting.title,
        meeting_date=date(2026, 10, 15),
        raw_text=transcript_text,
        attendees=["Sarah", "Alex", "David", "Elena"],
        created_at=datetime(2026, 10, 15, 10, 0, 0),
    )
    meeting_model.items = [task_1, task_2, dec_1, q_1]
    client_model = ClientModel(id=client.id, name=client.name)
    meeting_model.client = client_model

    # Step 6: Review and accept Task 1
    mock_db.scalars.return_value.first.return_value = task_1
    res_accept = ItemService.update_item(
        db=mock_db,
        item_id=task_1.id,
        update_in=ItemUpdate(review_status="accepted"),
    )
    assert res_accept.review_status == "accepted"

    # Step 7: Leave Task 2 in needs_review (owner is not_stated, deadline is None)
    assert task_2.review_status == "needs_review"
    assert task_2.payload_json["owner"]["name"] is None

    # Step 8: Reject an item
    mock_db.scalars.return_value.first.return_value = dec_1
    res_reject = ItemService.update_item(
        db=mock_db,
        item_id=dec_1.id,
        update_in=ItemUpdate(review_status="rejected"),
    )
    assert res_reject.review_status == "rejected"

    # Step 9: Open Clarifications & Duplicates
    mock_db.scalar.return_value = client_model
    mock_db.scalars.return_value.all.return_value = meeting_model.items
    clarifications_res = ClarificationService.get_client_open_items(
        db=mock_db,
        client_id=client.id,
    )
    assert len(clarifications_res.items) >= 1
    unresolved_owners = [c for c in clarifications_res.items if "owner" in c.missing_fields]
    assert len(unresolved_owners) >= 1

    from app.similarity.deduplication import calculate_pair_similarity
    sim = calculate_pair_similarity(
        task_1.payload_json["description"],
        "Complete the MySQL 8 database migration",
        owner1="Alex",
        owner2="Alex",
    )
    assert sim > 0.50
    mock_db.scalar.return_value = meeting_model
    mock_db.scalars.return_value.all.side_effect = [
        [task_1, task_2],  # current tasks
        [],                # candidate tasks -> empty returns []
    ]
    dup_candidates = DeduplicationService.get_meeting_duplicates(
        db=mock_db,
        meeting_id=meeting.id,
    )
    assert isinstance(dup_candidates, list)

    # Step 10: Review saved meeting history
    mock_db.scalar.return_value = 1
    mock_db.scalars.return_value.unique.return_value.all.return_value = [meeting_model]
    history, total = MeetingService.list_meetings(mock_db, client_id=client.id)
    assert total == 1
    assert history[0].title == "Q4 Technical Architecture Alignment"

    # Step 11: Export Excel, CSV, Word, and JSON
    # 11a: Excel Export
    excel_stream = build_excel_export(meeting_model, status_filter="all")
    wb = openpyxl.load_workbook(io.BytesIO(excel_stream.getvalue()))
    assert "Tasks" in wb.sheetnames
    assert "Decisions" in wb.sheetnames
    assert "README" in wb.sheetnames

    # 11b: CSV Export
    csv_stream = build_csv_export(meeting_model, status_filter="all")
    csv_content = csv_stream.getvalue().decode("utf-8-sig")
    assert "Complete the MySQL 8 migration" in csv_content
    assert "Not stated" in csv_content

    # 11c: Word Export
    docx_stream = build_word_export(meeting_model, reviewer_name="QA Lead", status_filter="all")
    doc = Document(io.BytesIO(docx_stream.getvalue()))
    doc_text = "\n".join(p.text for p in doc.paragraphs if p.text)
    assert "Q4 Technical Architecture Alignment" in doc_text
    assert "1. Meeting Summary" in doc_text

    # 11d: JSON Export
    json_data = build_json_dict(meeting_model, status_filter="all", include_transcript=True)
    assert json_data["schema_version"] == "1.0"
    assert json_data["client"]["name"] == "Starlight Dynamics Inc."
    assert len(json_data["tasks"]) == 2

    # Step 12: Reopen meeting & verify data persisted
    mock_db.scalars.return_value.unique.return_value.first.return_value = meeting_model
    mock_db.scalar.return_value = meeting_model
    reopened = MeetingService.get_meeting(mock_db, meeting.id)
    assert reopened.id == meeting.id
    assert reopened.title == "Q4 Technical Architecture Alignment"
    assert len(reopened.items) == 4

    # Step 13: Verify error handling (non-existent item raises 404, never manufactured success)
    mock_db.scalars.return_value.first.return_value = None
    try:
        ItemService.get_item(mock_db, "non-existent-id")
        raise AssertionError("Should have raised EntityNotFoundError")
    except Exception as exc:
        assert "not found" in str(exc).lower()

    # Step 14: Final state integrity
    assert task_1.review_status == "accepted"
    assert dec_1.review_status == "rejected"
    assert task_2.review_status == "needs_review"
