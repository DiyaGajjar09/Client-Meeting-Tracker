"""Tests for export generation (Excel, CSV, Word, JSON) and export API endpoints."""

import io
from datetime import date, datetime
from unittest.mock import MagicMock

import openpyxl
from docx import Document
from fastapi.testclient import TestClient

from app.db.models import ClientModel, ItemModel, MeetingModel
from app.db.session import get_db
from app.export.common import (
    format_date_display,
    generate_safe_export_filename,
    sanitize_spreadsheet_value,
)
from app.export.csv_export import build_csv_export
from app.export.excel import build_excel_export
from app.export.json_export import build_json_dict
from app.export.schemas import ColumnMappingItem
from app.export.word import build_word_export
from app.main import app


def create_sample_meeting_and_items():
    client = ClientModel(id="client-test-1", name="Acme Corporation")
    meeting = MeetingModel(
        id="meeting-test-1",
        client_id="client-test-1",
        title="Q4 Product Review",
        meeting_date=date(2026, 10, 15),
        created_at=datetime(2026, 10, 15, 10, 0, 0),
        raw_text="Alice: Finalize migration by Nov 15.\nBob: Look at vendor API rate limits.",
    )
    meeting.client = client

    items = [
        ItemModel(
            id="item-task-1",
            meeting_id="meeting-test-1",
            item_type="task",
            payload_json={
                "description": "Finalize the data migration plan",
                "owner": {"name": "Tom Hargreaves", "state": "explicit"},
                "due": {"resolved_date": "2026-11-15", "raw_text": "15 Nov 2026"},
                "priority": "high",
                "confidence": "high",
                "evidence": [
                    {
                        "quote": "Tom will finalize the data migration plan by 15 Nov 2026.",
                        "line_start": 1,
                        "line_end": 2,
                    }
                ],
            },
            review_status="accepted",
        ),
        ItemModel(
            id="item-task-2",
            meeting_id="meeting-test-1",
            item_type="task",
            payload_json={
                "description": "=cmd|' /C calc'!A0 Formula Test Task",  # Injection attempt
                "owner": {"name": None, "state": "not_stated"},  # Missing owner
                "due": {"resolved_date": None, "raw_text": "this week"},  # Unresolved deadline
                "priority": "medium",
                "confidence": "medium",
                "evidence": [
                    {
                        "quote": "Someone should check the formula.",
                        "line_start": 3,
                        "line_end": 4,
                    }
                ],
                "clarification_needed": ["Assign an owner for formula check"],
            },
            review_status="needs_review",
        ),
        ItemModel(
            id="item-dec-1",
            meeting_id="meeting-test-1",
            item_type="decision",
            payload_json={
                "description": "Use FastAPI and MySQL for backend architecture",
                "owner": {"name": "Alice", "state": "explicit"},
                "confidence": "high",
                "evidence": [{"quote": "Team agreed on FastAPI and MySQL."}],
            },
            review_status="accepted",
        ),
        ItemModel(
            id="item-risk-1",
            meeting_id="meeting-test-1",
            item_type="risk",
            payload_json={
                "description": "Vendor API rate limit may throttle production jobs",
                "severity": "high",
                "confidence": "high",
                "evidence": [{"quote": "Vendor rate limit is only 60 req/min."}],
            },
            review_status="accepted",
        ),
        ItemModel(
            id="item-q-1",
            meeting_id="meeting-test-1",
            item_type="open_question",
            payload_json={
                "question": "Will the client provide test credentials by Monday?",
                "directed_to": "Acme Corp",
                "confidence": "medium",
                "evidence": [{"quote": "Can they provide staging keys?"}],
            },
            review_status="needs_review",
        ),
    ]

    meeting.items = items
    return meeting, items


def test_sanitize_spreadsheet_value():
    """Verify formula injection characters (=, +, -, @, \\t, \\r) are prefixed with '."""
    assert sanitize_spreadsheet_value("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_spreadsheet_value("+12345") == "'+12345"
    assert sanitize_spreadsheet_value("-cmd") == "'-cmd"
    assert sanitize_spreadsheet_value("@google.com") == "'@google.com"
    assert sanitize_spreadsheet_value("\tTabbed") == "'\tTabbed"
    assert sanitize_spreadsheet_value("\rCarriage") == "'\rCarriage"
    assert sanitize_spreadsheet_value("Normal text") == "Normal text"
    assert sanitize_spreadsheet_value(None) == ""
    assert sanitize_spreadsheet_value(123) == 123


def test_format_date_display():
    assert format_date_display(date(2026, 10, 15)) == "15-Oct-2026"
    assert format_date_display("2026-10-15") == "15-Oct-2026"
    assert format_date_display(None) == ""


def test_generate_safe_export_filename():
    filename = generate_safe_export_filename("Acme / Corp * Special", "Q4 : Review?", date(2026, 10, 15), "xlsx")
    assert filename == "Acme___Corp___Special_Q4___Review_2026-10-15.xlsx"


def test_excel_export_generation():
    """Verify openpyxl workbook structure with 5 sheets and formatting."""
    meeting, items = create_sample_meeting_and_items()

    excel_stream = build_excel_export(meeting, status_filter="all")
    excel_bytes = excel_stream.getvalue()
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert sheet_names == ["Tasks", "Decisions", "Risks", "Open Questions", "README"]

    # Verify Tasks sheet
    ws_tasks = wb["Tasks"]
    assert ws_tasks.freeze_panes == "A2"
    headers = [cell.value for cell in ws_tasks[1]]
    assert "Task ID" in headers
    assert "Task" in headers
    assert "Owner" in headers
    assert "Due Date" in headers

    # Row 2 (Task 1): Tom Hargreaves, 15-Nov-2026
    assert ws_tasks.cell(row=2, column=5).value == "Tom Hargreaves"
    assert ws_tasks.cell(row=2, column=7).value == "15-Nov-2026"

    # Row 3 (Task 2): Formula injection prevented & Not stated owner
    assert ws_tasks.cell(row=3, column=4).value.startswith("'=")
    assert ws_tasks.cell(row=3, column=5).value == "Not stated"
    assert ws_tasks.cell(row=3, column=7).value in (None, "")

    # Verify README sheet
    ws_readme = wb["README"]
    assert "EXPORT DOCUMENTATION" in ws_readme.cell(row=1, column=1).value


def test_csv_export_generation_and_column_mapping():
    """Verify CSV generation with UTF-8 BOM, formula defense, and column mapping."""
    meeting, items = create_sample_meeting_and_items()

    # Custom mapping: Task, Owner, Status only
    custom_cols = [
        ColumnMappingItem(id="description", label="Action Item", selected=True),
        ColumnMappingItem(id="owner", label="Assignee", selected=True),
        ColumnMappingItem(id="status", label="Current Status", selected=True),
        ColumnMappingItem(id="due_date", label="Due Date", selected=False),
    ]

    csv_stream = build_csv_export(meeting, column_mapping=custom_cols, status_filter="all")
    csv_bytes = csv_stream.getvalue()

    # Must start with UTF-8 BOM
    assert csv_bytes.startswith(b"\xef\xbb\xbf")

    decoded = csv_bytes.decode("utf-8-sig")
    lines = decoded.strip().splitlines()
    assert lines[0] == "Action Item,Assignee,Current Status"

    # Task 1 row
    assert "Finalize the data migration plan,Tom Hargreaves,accepted" in lines[1]
    # Task 2 row (formula defense applied, missing owner is Not stated)
    assert "'=cmd|' /C calc'!A0 Formula Test Task,Not stated,needs_review" in lines[2]


def test_word_export_generation():
    """Verify python-docx meeting minutes with title, all 6 sections, and footer."""
    meeting, items = create_sample_meeting_and_items()

    docx_stream = build_word_export(meeting, reviewer_name="Jane Doe (Lead QA)", status_filter="all")
    docx_bytes = docx_stream.getvalue()
    assert len(docx_bytes) > 0

    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    full_text = "\n".join(paragraphs)

    assert "Q4 Product Review" in full_text
    assert "Meeting Minutes" in full_text
    assert "1. Meeting Summary" in full_text
    assert "2. Formal Decisions" in full_text
    assert "3. Action Items" in full_text
    assert "4. Risks & Mitigations" in full_text
    assert "5. Open Questions" in full_text
    assert "6. Clarifications Required" in full_text

    # Tables verified
    assert len(doc.tables) >= 2


def test_json_export_structure():
    """Verify structured JSON output adhering to schema 1.0 without fabricated data."""
    meeting, items = create_sample_meeting_and_items()

    data = build_json_dict(meeting, status_filter="all", include_transcript=True)
    assert data["schema_version"] == "1.0"
    assert data["client"]["name"] == "Acme Corporation"
    assert data["meeting"]["title"] == "Q4 Product Review"
    assert data["meeting"].get("raw_text") is not None
    assert len(data["tasks"]) == 2
    assert len(data["decisions"]) == 1
    assert len(data["risks"]) == 1
    assert len(data["open_questions"]) == 1

    # Missing owner strictly null in JSON, no invented owner
    task2 = [t for t in data["tasks"] if t["id"] == "item-task-2"][0]
    assert task2["owner"]["name"] is None
    assert task2["due_date"]["resolved_date"] is None
    assert task2["due_date"]["raw_text"] == "this week"

    # Without transcript
    data_no_t = build_json_dict(meeting, status_filter="all", include_transcript=False)
    assert "raw_text" not in data_no_t["meeting"]


def test_export_service_and_api_endpoint():
    """Verify FastAPI POST /meetings/{meeting_id}/export endpoint."""
    meeting, items = create_sample_meeting_and_items()
    mock_db = MagicMock()
    mock_db.scalar.return_value = meeting

    def override_get_db():
        yield mock_db

    from app.core.auth import require_auth
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[require_auth] = lambda: {"sub": "test_user"}
    client = TestClient(app)

    try:
        # 1. Export Excel
        res_xlsx = client.post(
            "/meetings/meeting-test-1/export",
            json={"format": "xlsx", "status_filter": "all"},
        )
        assert res_xlsx.status_code == 200
        assert (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            in res_xlsx.headers["content-type"]
        )
        assert "Content-Disposition" in res_xlsx.headers
        assert ".xlsx" in res_xlsx.headers["Content-Disposition"]

        # 2. Export CSV
        res_csv = client.post(
            "/meetings/meeting-test-1/export",
            json={"format": "csv"},
        )
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers["content-type"]
        assert ".csv" in res_csv.headers["Content-Disposition"]

        # 3. Export Word
        res_docx = client.post(
            "/meetings/meeting-test-1/export",
            json={"format": "docx"},
        )
        assert res_docx.status_code == 200
        assert (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            in res_docx.headers["content-type"]
        )
        assert ".docx" in res_docx.headers["Content-Disposition"]

        # 4. Export JSON
        res_json = client.post(
            "/meetings/meeting-test-1/export",
            json={"format": "json"},
        )
        assert res_json.status_code == 200
        assert "application/json" in res_json.headers["content-type"]
        body = res_json.json()
        assert body["schema_version"] == "1.0"

        # 5. Invalid format 422
        res_invalid = client.post(
            "/meetings/meeting-test-1/export",
            json={"format": "pdf"},
        )
        assert res_invalid.status_code == 422

        # 6. Meeting not found 404
        mock_db.scalar.return_value = None
        res_404 = client.post(
            "/meetings/non-existent/export",
            json={"format": "xlsx"},
        )
        assert res_404.status_code == 404

    finally:
        app.dependency_overrides.pop(get_db, None)
