from __future__ import annotations

import io
from datetime import UTC, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.db.models import ItemModel, MeetingModel
from app.export.common import categorize_meeting_items, format_date_display


def set_cell_background(cell, hex_color: str):
    """Set background shading color of a docx table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding inside a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def build_word_export(
    meeting: MeetingModel,
    reviewer_name: str | None = None,
    status_filter: str = "all",
) -> io.BytesIO:
    """Generate professional Word meeting minutes (.docx) adhering to the PRD layout."""
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run(
            f"Generated from meeting notes and reviewed by {reviewer_name or 'Designated Reviewer'}. "
            f"({datetime.now(UTC).strftime('%d-%b-%Y')})"
        )
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    client_name = meeting.client.name if meeting.client else "Client Organization"
    meeting_title = meeting.title
    meeting_date = format_date_display(meeting.meeting_date) or "Date Unspecified"
    attendees_list = meeting.attendees or []
    attendees_str = ", ".join(attendees_list) if attendees_list else "None recorded"

    items_by_cat = categorize_meeting_items(meeting.items or [], status_filter=status_filter)
    tasks = items_by_cat["tasks"]
    decisions = items_by_cat["decisions"]
    risks = items_by_cat["risks"]
    questions = items_by_cat["open_questions"]

    # -------------------------------------------------------------
    # Title Block
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(meeting_title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Meeting Minutes — {client_name}")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    sub_p.paragraph_format.space_after = Pt(12)

    # Metadata Panel Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_cells = [
        (0, 0, "Meeting Date:", meeting_date),
        (0, 1, "Generated On:", datetime.now(UTC).strftime("%d-%b-%Y %H:%M UTC")),
        (1, 0, "Client:", client_name),
        (1, 1, "Participants:", attendees_str),
    ]
    for r, c, label, val in meta_cells:
        cell = meta_table.cell(r, c)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        lbl_run = p.add_run(f"{label} ")
        lbl_run.bold = True
        lbl_run.font.size = Pt(9.5)
        lbl_run.font.color.rgb = RGBColor(71, 85, 105)
        val_run = p.add_run(val)
        val_run.font.size = Pt(9.5)
        val_run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # Section 1: Meeting Summary
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Meeting Summary", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    summary_text = (
        f"This session addressed project deliverables and commitments for {client_name}. "
        f"A total of {len(tasks)} tasks, {len(decisions)} decisions, {len(risks)} risks, and "
        f"{len(questions)} open questions were tracked and reviewed from the transcript."
    )
    sum_p = doc.add_paragraph(summary_text)
    sum_p.paragraph_format.line_spacing = 1.15
    sum_p.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # Section 2: Decisions
    # -------------------------------------------------------------
    h2 = doc.add_heading("2. Formal Decisions", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    if not decisions:
        p = doc.add_paragraph("No formal decisions were agreed upon during this meeting.")
        p.runs[0].font.italic = True
        p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    else:
        for idx, item in enumerate(decisions, start=1):
            payload = item.payload_json or {}
            dec_stmt = payload.get("statement", payload.get("description", ""))
            dec_by = payload.get("decided_by", "Consensus")
            ev_list = payload.get("evidence", [])
            ev_quote = ev_list[0].get("quote", "") if ev_list and isinstance(ev_list[0], dict) else ""

            dp = doc.add_paragraph(style="List Bullet")
            r_stmt = dp.add_run(f"Decision {idx}: {dec_stmt} ")
            r_stmt.bold = True
            r_by = dp.add_run(f"(Decided by: {dec_by} · Status: {item.review_status.upper()})")
            r_by.font.size = Pt(9)
            r_by.font.color.rgb = RGBColor(71, 85, 105)

            if ev_quote:
                ep = doc.add_paragraph()
                ep.paragraph_format.left_indent = Inches(0.4)
                ep.paragraph_format.space_after = Pt(4)
                e_run = ep.add_run(f'"{ev_quote}"')
                e_run.font.italic = True
                e_run.font.size = Pt(9)
                e_run.font.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------
    # Section 3: Action Items
    # -------------------------------------------------------------
    h3 = doc.add_heading("3. Action Items", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    accepted_tasks = [t for t in tasks if t.review_status == "accepted"]
    unresolved_tasks = [t for t in tasks if t.review_status != "accepted"]

    def render_task_table(task_list: list[ItemModel], table_title: str):
        if not task_list:
            p = doc.add_paragraph(f"No {table_title.lower()} recorded.")
            p.runs[0].font.italic = True
            p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
            return

        doc.add_paragraph(table_title, style="List Bullet").runs[0].bold = True

        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ["Task Description", "Owner", "Due Date", "Status", "Confidence"]
        col_widths = [Inches(2.8), Inches(1.2), Inches(1.1), Inches(1.0), Inches(0.9)]

        hdr_cells = table.rows[0].cells
        for i, h_text in enumerate(headers):
            hdr_cells[i].text = h_text
            hdr_cells[i].width = col_widths[i]
            set_cell_background(hdr_cells[i], "1E293B")
            set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
            p = hdr_cells[i].paragraphs[0]
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.size = Pt(9)

        for item in task_list:
            payload = item.payload_json or {}
            owner_info = payload.get("owner", {})
            owner_name = owner_info.get("name") or "Not stated"

            due_info = payload.get("due", {})
            resolved_date = due_info.get("resolved_date")
            raw_wording = due_info.get("raw_text")
            date_display = format_date_display(resolved_date) if resolved_date else (raw_wording or "Not stated")

            row_cells = table.add_row().cells
            row_data = [
                payload.get("description", ""),
                owner_name,
                date_display,
                item.review_status.upper(),
                payload.get("confidence", "medium").capitalize(),
            ]
            for i, val in enumerate(row_data):
                row_cells[i].text = val
                set_cell_margins(row_cells[i], top=80, bottom=80, left=120, right=120)
                p = row_cells[i].paragraphs[0]
                p.runs[0].font.size = Pt(8.5)
                p.runs[0].font.name = "Arial"
                if i == 1 and owner_name == "Not stated":
                    set_cell_background(row_cells[i], "FEF3C7")
                    p.runs[0].font.italic = True
                if i == 2 and not resolved_date:
                    set_cell_background(row_cells[i], "FEF3C7")

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    render_task_table(accepted_tasks, "Approved Action Items for Implementation")
    if unresolved_tasks:
        render_task_table(unresolved_tasks, "Action Items Awaiting Final Review or Clarification")

    # -------------------------------------------------------------
    # Section 4: Risks
    # -------------------------------------------------------------
    h4 = doc.add_heading("4. Risks & Mitigations", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    if not risks:
        p = doc.add_paragraph("No outstanding project risks were identified in this meeting.")
        p.runs[0].font.italic = True
        p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    else:
        for idx, item in enumerate(risks, start=1):
            payload = item.payload_json or {}
            owner_info = payload.get("owner", {})
            rp = doc.add_paragraph(style="List Bullet")
            rp.add_run(f"Risk {idx}: {payload.get('description', '')}\n").bold = True
            rp.add_run(f"• Impact: {payload.get('impact', 'Unspecified')}\n")
            rp.add_run(f"• Mitigation: {payload.get('mitigation', 'Pending formulation')}\n")
            rp.add_run(f"• Owner: {owner_info.get('name') or 'Not stated'} · Status: {item.review_status}")

    # -------------------------------------------------------------
    # Section 5: Open Questions
    # -------------------------------------------------------------
    h5 = doc.add_heading("5. Open Questions", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    if not questions:
        p = doc.add_paragraph("No unanswered open questions remain from this meeting.")
        p.runs[0].font.italic = True
        p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    else:
        for idx, item in enumerate(questions, start=1):
            payload = item.payload_json or {}
            owner_info = payload.get("owner", {})
            qp = doc.add_paragraph(style="List Bullet")
            qp.add_run(f"Question {idx}: {payload.get('question', payload.get('description', ''))}\n").bold = True
            qp.add_run(f"• Directed To: {owner_info.get('name') or 'Unassigned'} · Review Status: {item.review_status}")

    # -------------------------------------------------------------
    # Section 6: Clarifications Required
    # -------------------------------------------------------------
    h6 = doc.add_heading("6. Clarifications Required", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)

    missing_owner_count = sum(
        1 for t in tasks if (t.payload_json or {}).get("owner", {}).get("state") == "not_stated"
    )
    missing_due_count = sum(
        1 for t in tasks if not (t.payload_json or {}).get("due", {}).get("resolved_date")
    )
    low_conf_count = sum(
        1 for it in (meeting.items or []) if (it.payload_json or {}).get("confidence") == "low"
    )

    clarification_summary = [
        f"• Tasks missing assigned owners: {missing_owner_count}",
        f"• Tasks with unconfirmed or vague deadlines: {missing_due_count}",
        f"• Items requiring manual verification due to low confidence: {low_conf_count}",
        f"• Total open questions requiring stakeholder feedback: {len(questions)}",
    ]
    for c_line in clarification_summary:
        doc.add_paragraph(c_line)

    # Save to memory buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
