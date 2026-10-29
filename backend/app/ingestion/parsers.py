import io
import re

from app.core.config import settings
from app.core.exceptions import (
    InvalidInputError,
    PayloadTooLargeError,
    UnsupportedMediaTypeError,
)

SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf", ".vtt", ".srt"}


def extract_text_from_txt(content: bytes) -> str:
    """Decode raw text bytes trying UTF-8, UTF-8 with BOM, and Latin-1 fallbacks."""
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise InvalidInputError("Unable to decode text file. Ensure it is encoded in valid UTF-8.")


def extract_text_from_docx(content: bytes) -> str:
    """Extract paragraphs and table text from DOCX file using python-docx with ZIP bomb protection."""
    try:
        import docx
    except ImportError:
        raise InvalidInputError("python-docx is not installed on the system.")

    # Application Security: Inspect ZIP archive entries to defend against decompression bombs and directory traversal
    try:
        import zipfile

        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            total_uncompressed = 0
            for info in zf.infolist():
                if ".." in info.filename or info.filename.startswith(("/", "\\")):
                    raise InvalidInputError("DOCX archive contains unsafe path traversal entries.")
                total_uncompressed += info.file_size

            # Maximum uncompressed threshold: 50MB
            max_uncompressed_limit = 50 * 1024 * 1024
            if total_uncompressed > max_uncompressed_limit:
                raise PayloadTooLargeError(
                    f"DOCX expanded size ({total_uncompressed // 1048576} MB) exceeds maximum safe expansion limit."
                )

            # Expansion ratio check: reject archives with > 50:1 ratio
            if len(content) > 100 and (total_uncompressed / max(len(content), 1)) > 50:
                raise PayloadTooLargeError(
                    "DOCX archive exhibits an abnormal compression ratio indicative of a decompression bomb."
                )
    except zipfile.BadZipFile:
        raise InvalidInputError("Corrupt or invalid DOCX archive structure.")
    except (InvalidInputError, PayloadTooLargeError):
        raise
    except Exception as exc:
        raise InvalidInputError(f"Error inspecting DOCX package: {exc!s}")

    try:
        doc = docx.Document(io.BytesIO(content))
        lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                lines.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    lines.append(" | ".join(row_cells))
        extracted = "\n".join(lines).strip()
        if not extracted:
            raise InvalidInputError("DOCX file does not contain any readable text.")
        return extracted
    except Exception as exc:
        if isinstance(exc, (InvalidInputError, PayloadTooLargeError)):
            raise
        raise InvalidInputError(f"Corrupt or unreadable DOCX document: {exc!s}")


def extract_text_from_pdf(content: bytes) -> str:
    """Extract selectable text from PDF using pypdf. Scanned/OCR PDFs are not supported in v1.0."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise InvalidInputError("pypdf is not installed on the system.")

    try:
        reader = PdfReader(io.BytesIO(content))
        lines = []
        for _idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            stripped = page_text.strip()
            if stripped:
                lines.append(stripped)

        extracted = "\n".join(lines).strip()
        if not extracted:
            raise InvalidInputError(
                "PDF contains no selectable text. Scanned PDFs requiring OCR are not supported in v1.0."
            )
        return extracted
    except Exception as exc:
        if isinstance(exc, InvalidInputError):
            raise
        raise InvalidInputError(f"Corrupt or unreadable PDF document: {exc!s}")


def extract_text_from_vtt(content: bytes) -> str:
    """Parse WebVTT subtitle/transcript text, removing timestamps while preserving speaker turns."""
    raw = extract_text_from_txt(content)
    lines = []
    # Strip WEBVTT header and cue timing lines
    for line in raw.splitlines():
        line = line.strip()
        if not line or line.startswith(("WEBVTT", "NOTE")):
            continue
        # Remove timestamp lines (e.g. 00:01:20.000 --> 00:01:23.000)
        if "-->" in line:
            continue
        # Strip HTML voice/speaker tags like <v SpeakerName>
        cleaned = re.sub(r"<v\s+([^>]+)>", r"\1: ", line)
        cleaned = re.sub(r"</v>", "", cleaned)
        cleaned = re.sub(r"<[^>]+>", "", cleaned).strip()
        if cleaned:
            lines.append(cleaned)

    extracted = "\n".join(lines).strip()
    if not extracted:
        raise InvalidInputError("VTT file contains no valid dialogue text.")
    return extracted


def extract_text_from_srt(content: bytes) -> str:
    """Parse SubRip (SRT) subtitle/transcript text, removing index numbers and timestamps."""
    raw = extract_text_from_txt(content)
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        if not line or line.isdigit():
            continue
        # Remove timestamp lines (00:00:20,000 --> 00:00:23,000)
        if "-->" in line:
            continue
        cleaned = re.sub(r"<[^>]+>", "", line).strip()
        if cleaned:
            lines.append(cleaned)

    extracted = "\n".join(lines).strip()
    if not extracted:
        raise InvalidInputError("SRT file contains no valid dialogue text.")
    return extracted


def parse_transcript_file(filename: str, content: bytes) -> tuple[str, str]:
    """Validate and parse an uploaded transcript file into clean text.

    Returns:
        Tuple of (clean_text, filename)
    """
    if len(content) > settings.MAX_UPLOAD_BYTES:
        max_mb = settings.MAX_UPLOAD_BYTES // (1024 * 1024)
        raise PayloadTooLargeError(f"Uploaded file exceeds maximum allowed limit of {max_mb} MB.")

    if len(content) == 0:
        raise InvalidInputError("Uploaded file is empty.")

    # Detect extension
    dot_index = filename.rfind(".")
    if dot_index == -1:
        raise UnsupportedMediaTypeError("Uploaded file has no file extension.")

    ext = filename[dot_index:].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedMediaTypeError(
            f"File extension '{ext}' is unsupported. Allowed formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if ext == ".txt":
        text = extract_text_from_txt(content)
    elif ext == ".docx":
        text = extract_text_from_docx(content)
    elif ext == ".pdf":
        text = extract_text_from_pdf(content)
    elif ext == ".vtt":
        text = extract_text_from_vtt(content)
    elif ext == ".srt":
        text = extract_text_from_srt(content)
    else:
        raise UnsupportedMediaTypeError(f"Handler not found for extension {ext}")

    cleaned = text.strip()
    if not cleaned:
        raise InvalidInputError("Uploaded transcript contains no readable text content.")

    return cleaned, filename
